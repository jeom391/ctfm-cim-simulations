"""Downloadable numeric tables and plots; strings never become spreadsheet formulas."""
import csv
import io
import json
from pathlib import Path
import math

def scalar(value):
    return json.dumps(value,ensure_ascii=False,allow_nan=False) if isinstance(value,(dict,list)) else value

def safe_csv(value):
    value=scalar(value)
    if isinstance(value,str) and value.lstrip().startswith(("=","+","-","@")):
        return "'"+value
    return value

def export_result(result, output_dir):
    from openpyxl import Workbook
    output_dir=Path(output_dir)
    (output_dir/"summary.json").write_text(json.dumps(result,ensure_ascii=False,indent=2,allow_nan=False),encoding="utf-8")
    tables=dict(result.get("tables",{}))
    if result.get("runs"):
        tables["runs"]=result["runs"]
    tables["exclusions"]=result.get("exclusions",[])
    workbook=Workbook();workbook.remove(workbook.active)
    for name,rows in tables.items():
        if not rows: continue
        columns=list(dict.fromkeys(key for row in rows for key in row))
        # Core owns names; prevent accidentally constructing a file path from data.
        name="".join(c for c in name if c.isalnum() or c in "_-")[:40] or "table"
        with (output_dir/(name+".csv")).open("w",encoding="utf-8-sig",newline="") as stream:
            writer=csv.writer(stream);writer.writerow(columns)
            writer.writerows([[safe_csv(row.get(key)) for key in columns] for row in rows])
        sheet=workbook.create_sheet(name[:31])
        # sheet.max_row scans every cell, which made this loop quadratic on 24k-row tables.
        for index,values in enumerate([columns]+[[scalar(row.get(key)) for key in columns] for row in rows],start=1):
            sheet.append(values)
            for column,value in enumerate(values,start=1):
                if isinstance(value,str): sheet.cell(index,column).data_type="s"
        sheet.freeze_panes="A2";sheet.auto_filter.ref=sheet.dimensions
    if not workbook.worksheets:
        sheet=workbook.create_sheet("summary");sheet.append(["status","No tabular rows"])
    workbook.save(output_dir/"tables.xlsx")
    plot_result(result,output_dir)
    lines=["# Computation result","", "Measurement evidence and modeling assumptions are recorded in summary.json.",""]
    if result.get("summary"):lines.append(json.dumps(result["summary"],ensure_ascii=False,indent=2,allow_nan=False))
    for warning in result.get("warnings",[]):lines.append("- "+warning)
    (output_dir/"summary.md").write_text("\n".join(lines)+"\n",encoding="utf-8")

def plot_result(result, output_dir):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig,ax=plt.subplots(figsize=(8,4.5),layout="constrained")
    kind=result.get("kind")
    rows=result.get("tables",{}).get("raw",[])
    if kind=="pulse_states":
        for direction in ("ltp","ltd"):
            states=[s for s in result["states"] if s["direction"]==direction]
            ax.plot([s["extraction_index"] for s in states],[s["conductance_s"] for s in states],"o-",label=direction)
        ax.set(xlabel="Observed extraction index",ylabel="Absolute conductance (S)",title="Observed pulse states")
    elif kind in ("iv","d2d"):
        for file_id in dict.fromkeys(row["file_id"] for row in rows):
            series=[r for r in rows if r["file_id"]==file_id]
            ax.plot([r["vgs_v"] for r in series],[r["id_a"] for r in series],label=file_id[:8])
        ax.set(xlabel="VGS (V)",ylabel="ID (A)",title="Measured IV curves")
    elif kind=="retention":
        for direction in ("program","erase"):
            series=[r for r in result["tables"].get("retention_fit",[]) if r["direction"]==direction]
            ax.scatter([r["time_s"] for r in series],[r["id_a"] for r in series],s=12,label=direction+" raw")
            series=sorted(series,key=lambda r:r["time_s"])
            ax.plot([r["time_s"] for r in series],[r["fit_id_a"] for r in series],label=direction+" fit")
        ax.set(xscale="log",xlabel="Time (s)",ylabel="Absolute ID (A)",title="Raw-current logarithmic retention fit")
    elif result.get("runs"):
        runs=[r for r in result["runs"] if r.get("kind")=="ALL" and r["status"]=="succeeded"]
        for candidate in dict.fromkeys(r["candidate_id"] for r in runs):
            subset=[r for r in runs if r["candidate_id"]==candidate]
            years=sorted({r["years"] for r in subset})
            ax.plot(years,[sum(r["accuracy"] for r in subset if r["years"]==y)/sum(r["years"]==y for r in subset) for y in years],"o-",label=candidate)
        ax.set(xlabel="Years (extrapolation when beyond measured interval)",ylabel="Accuracy (fraction)",title="Valid inference runs only",ylim=(0,1))
    else:
        plt.close(fig);return
    handles,labels=ax.get_legend_handles_labels()
    if handles:ax.legend(fontsize=8)
    ax.grid(alpha=.2)
    fig.savefig(output_dir/"plot.png",dpi=160);plt.close(fig)
