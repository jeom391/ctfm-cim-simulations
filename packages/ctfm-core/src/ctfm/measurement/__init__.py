"""Explicit, provenance-preserving measurement analysis in SI units."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
from copy import deepcopy
from pathlib import Path
from statistics import mean, stdev

PARSER_VERSION = '1.0.0'
RETENTION_SOURCES = {'A1':('R1',0.),'A2':('R2',0.),'A3':('R3(1)',.5),'A4':('R4(1)',.1),'A5':('R5(1)',.5)}
MAX_ROWS = 1_000_000
MAX_COLUMNS = 256
MAX_CELLS = 5_000_000

def _bounded_rows(rows):
    total = 0
    for index, row in enumerate(rows, 1):
        total += len(row)
        if index > MAX_ROWS or len(row) > MAX_COLUMNS or total > MAX_CELLS:
            raise ValueError('Table parser row/column/cell limit exceeded')
        yield row

DEFAULTS = dict(read_tolerance_v=0.05, write_threshold_v=5.0, start_time_s=6.0,
                sample_offset_rows=2, iref_a=1e-6, vds_v=0.1, read_vgs_v=0.0,
                retention_start_time_s=10.0, crossing_segments={}, excluded_conditions={})
UNIT_FACTORS = {'time_s': {'s':1., 'ms':1e-3, 'us':1e-6, '쨉s':1e-6},
                'vgs_v': {'V':1., 'mV':1e-3},
                'id_a': {'A':1., 'mA':1e-3, 'uA':1e-6, '쨉A':1e-6, '關A':1e-6, 'nA':1e-9}}


def parse_table(data: bytes, filename: str, sheet: str | None = None) -> dict:
    """Read raw cells; identify a header only, never infer a calculation mapping."""
    if Path(filename).name.startswith('~$'):
        raise ValueError('Temporary spreadsheet files are excluded')
    extension = Path(filename).suffix.lower()
    sheets, warnings = [], []
    if extension == '.xls':
        raise ValueError('Legacy .xls is unsupported; convert to XLSX')
    if extension == '.csv':
        for encoding in ('utf-8-sig', 'cp949'):
            try:
                decoded = data.decode(encoding, errors='strict')
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError('CSV must be UTF-8 or CP949')
        try:
            raw = list(_bounded_rows(csv.reader(io.StringIO(decoded), strict=True)))
        except csv.Error as exc:
            raise ValueError(f'Invalid CSV: {exc}') from exc
        if encoding == 'cp949':
            warnings.append('CSV decoded using CP949')
    elif extension == '.xlsx':
        from openpyxl import load_workbook
        try:
            formulas = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
            cached = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        except Exception as exc:
            raise ValueError('Invalid XLSX workbook') from exc
        try:
            sheets = formulas.sheetnames
            if sheet is None:
                if len(sheets) != 1:
                    return dict(sheets=sheets, columns=[], rows=[], source_rows=[], warnings=['Select a worksheet explicitly'])
                sheet = sheets[0]
            if sheet not in sheets:
                raise ValueError('Selected worksheet does not exist')
            height = formulas[sheet].max_row or 0
            width = formulas[sheet].max_column or 0
            if height > MAX_ROWS or width > MAX_COLUMNS or height * width > MAX_CELLS:
                raise ValueError("Worksheet dimension exceeds parser limit")
            raw = []
            for formula_row, value_row in zip(_bounded_rows(formulas[sheet].iter_rows()), _bounded_rows(cached[sheet].iter_rows())):
                values = []
                for formula, value in zip(formula_row, value_row):
                    if formula.data_type == 'f' and value.value is None:
                        raise ValueError(f'Formula at {sheet}!{formula.coordinate} has no cached value; resave workbook or upload raw values')
                    if value.data_type == 'e':
                        raise ValueError(f'Excel error at {sheet}!{value.coordinate}')
                    v = value.value
                    if v is not None and not isinstance(v, (str, int, float, bool)):
                        v = str(v)
                    values.append(v)
                raw.append(values)
        finally:
            formulas.close(); cached.close()
    else:
        raise ValueError('Only CSV and XLSX files are supported')
    while raw and all(v is None or v == '' for v in raw[-1]):
        raw.pop()
    if not raw:
        raise ValueError('The table is empty')
    header_index = None
    for index, row in enumerate(raw):
        nonempty = [v for v in row if v is not None and str(v).strip()]
        if len(nonempty) >= 2 and all(isinstance(v, str) and not _numeric(v) for v in nonempty):
            header_index = index
            break
    if header_index is None:
        raise ValueError('A header row with explicit column labels is required')
    header = raw[header_index]
    while header and (header[-1] is None or header[-1] == ''):
        header = header[:-1]
    columns = [str(v).strip() if v is not None else '' for v in header]
    if not all(columns) or len(set(columns)) != len(columns):
        raise ValueError('Column labels must be nonempty and unique')
    rows, source_rows = [], []
    for n, row in enumerate(raw[header_index + 1:], start=header_index + 2):
        if len(row) > len(columns) and any(v not in ('', None) for v in row[len(columns):]):
            raise ValueError(f'Row {n} has more cells than the header')
        rows.append(dict(zip(columns, list(row[:len(columns)]) + [None] * (len(columns)-len(row)))))
        source_rows.append(n)
    return dict(sheets=sheets, columns=columns, rows=rows, source_rows=source_rows, warnings=warnings)


def _numeric(value):
    try:
        float(value)
        return True
    except (TypeError, ValueError):
        return False


def _finite(value, label):
    if value is None or isinstance(value, bool) or isinstance(value, str) and not value.strip():
        raise ValueError(f'{label}: a finite number is required')
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f'{label}: a finite number is required') from exc
    if not math.isfinite(number):
        raise ValueError(f'{label}: a finite number is required')
    return number


def _prepare(dataset, keys):
    for key in ('file_id','sha256','filename','device_id','condition_id'):
        if not dataset.get(key):
            raise ValueError(f'Dataset requires {key}')
    if len(dataset['sha256']) != 64 or any(c not in '0123456789abcdef' for c in dataset['sha256']):
        raise ValueError('Dataset sha256 must be a lowercase SHA256 digest')
    rows = dataset.get('rows', [])
    indices = dataset.get('source_rows', [])
    if not rows or len(rows) != len(indices):
        raise ValueError('Nonempty rows and corresponding source_rows are required')
    if any(isinstance(n, bool) or not isinstance(n,int) or n < 1 for n in indices) or any(b <= a for a,b in zip(indices,indices[1:])):
        raise ValueError('source_rows must be increasing positive 1-based row numbers')
    mapping, units = dataset.get('column_mapping', {}), dataset.get('units', {})
    if any(k not in mapping or k not in units for k in keys):
        raise ValueError('Explicit column mapping and units are required for every input')
    if len(set(mapping[k] for k in keys)) != len(keys):
        raise ValueError('Mapped columns must be distinct')
    factors = {}
    for key in keys:
        family = 'id_a' if key.endswith('id_a') else key
        if units[key] not in UNIT_FACTORS[family]:
            raise ValueError(f'Unsupported unit for {key}: {units[key]}')
        factors[key] = UNIT_FACTORS[family][units[key]]
    converted = []
    for row, source_row in zip(rows, indices):
        converted.append(dict(source_row=source_row, **{k:_finite(row.get(mapping[k]),f'{dataset["file_id"]} row {source_row} {k}')*factors[k] for k in keys}))
    if 'time_s' in keys and any(b['time_s'] < a['time_s'] for a,b in zip(converted,converted[1:])):
        raise ValueError('Time reversal is not allowed; preserve acquisition order')
    if 'time_s' in keys and any(r['time_s'] < 0 for r in converted):
        raise ValueError('Negative measurement time is invalid')
    return converted


def crossing_selection_key(dataset):
    """Stable per-input key, including shared-file column and row selections."""
    amplitude=dataset.get('sweep_amplitude_v')
    identity=[dataset['file_id'],dataset['sha256'],dataset.get('sheet'),dataset['device_id'],
              dataset.get('branch'),float(amplitude) if amplitude is not None else None,
              dataset['source_rows'],sorted(dataset['column_mapping'].items())]
    encoded=json.dumps(identity,ensure_ascii=False,separators=(',',':'),allow_nan=False).encode('utf-8')
    return 'input:'+hashlib.sha256(encoded).hexdigest()


def _validate_crossing_keys(datasets,choices):
    if not isinstance(choices,dict): raise ValueError('crossing_segments must be an object')
    keys={crossing_selection_key(d) for d in datasets}
    for key in choices:
        if key in keys: continue
        matching=[d for d in datasets if d['file_id']==key]
        if len(matching)>1:
            raise ValueError('Ambiguous file_id crossing selection; use each input selection_key: '+', '.join(crossing_selection_key(d) for d in matching))
        if not matching: raise ValueError('Unknown crossing selection key: '+str(key))
    for d in datasets:
        if d['file_id'] in choices and crossing_selection_key(d) in choices:
            raise ValueError('Specify either selection_key or unique file_id, not both')


def _crossing_selection(d,result):
    choices=result['settings']['crossing_segments']
    return choices.get(crossing_selection_key(d),choices.get(d['file_id']))


def _source(d):
    fields = ('file_id','sha256','filename','sheet','device_id','condition_id','branch','sweep_amplitude_v','direction','source_label','read_vgs_v','vds_v','header_read_vgs_v','current_basis')
    return {**{k:deepcopy(d.get(k)) for k in fields}, 'columns':deepcopy(d['column_mapping']), 'units':deepcopy(d['units']), 'source_rows':deepcopy(d['source_rows']), 'selection_key':crossing_selection_key(d)}


def _excluded(result, d, reason, **extra):
    result['exclusions'].append(dict(file_id=d['file_id'], device_id=d['device_id'], reason=reason, **extra))


def analyze(kind: str, datasets: list[dict], settings: dict) -> dict:
    aliases = {'iv_vth':'iv', 'vth_mw':'iv', 'ltp_ltd':'pulse_states'}
    kind = aliases.get(kind, kind)
    keys = {'pulse_states':['time_s','id_a','vgs_v'], 'iv':['vgs_v','id_a'], 'd2d':['vgs_v','id_a'], 'retention':['time_s','program_id_a','erase_id_a']}
    if kind not in keys or not datasets:
        raise ValueError('Supported analysis kind and nonempty datasets are required')
    unknown = set(settings) - set(DEFAULTS)
    if unknown:
        raise ValueError(f'Unknown analysis settings: {sorted(unknown)}')
    config = {**deepcopy(DEFAULTS), **deepcopy(settings)}
    for key in ('read_tolerance_v','write_threshold_v','start_time_s','iref_a','vds_v','read_vgs_v','retention_start_time_s'):
        config[key] = _finite(config[key],key)
    if config['sample_offset_rows'] != 2 or isinstance(config['sample_offset_rows'],bool):
        raise ValueError('This parser version requires sample_offset_rows=2')
    if config['read_tolerance_v'] < 0 or config['write_threshold_v'] <= config['read_tolerance_v'] or config['start_time_s'] < 0:
        raise ValueError('Invalid pulse extraction thresholds')
    if not math.isclose(config['vds_v'],.1,abs_tol=1e-12) or config['iref_a'] <= 0 or config['read_vgs_v'] != 0 or config['retention_start_time_s'] < 10:
        raise ValueError('v1 requires VDS=0.1 V, read VGS=0 V, positive Iref, retention start >=10 s')
    if len({d.get('condition_id') for d in datasets}) != 1:
        raise ValueError('Do not mix A conditions in one analysis')
    prepared = [(d,_prepare(d,keys[kind])) for d in datasets]
    if kind in ('iv','d2d'): _validate_crossing_keys(datasets,config['crossing_segments'])
    result = dict(kind=kind, settings=config, condition_id=datasets[0]['condition_id'], summaries={}, tables={'raw':[]}, exclusions=[], warnings=[], provenance=[_source(d) for d in datasets])
    seen = set()
    for d, rows in prepared:
        signature=(d['sha256'],d.get('sheet'),tuple(d['source_rows']),tuple(sorted(d['column_mapping'].items())),d.get('branch'),d.get('direction'))
        if signature in seen:
            raise ValueError('Duplicate measurement dataset or physical-device input')
        seen.add(signature)
        for row, raw in zip(rows,d['rows']):
            result['tables']['raw'].append(dict(file_id=d['file_id'],**row,raw_cells=deepcopy(raw)))
    if kind == 'pulse_states': _pulse(prepared,result)
    elif kind == 'iv': _iv(prepared,result)
    elif kind == 'd2d': _d2d(prepared,result)
    else: _retention(prepared,result)
    # Trust boundary: no nonfinite values may leak into JSON artifacts.
    try:
        json.dumps(result,allow_nan=False)
    except (ValueError,TypeError) as exc:
        raise ValueError('Analysis contains non-finite or non-JSON input metadata') from exc
    return result


def _bias(d, expected_read=0):
    vds = _finite(d.get('vds_v'), 'vds_v')
    read = _finite(d.get('read_vgs_v'), 'read_vgs_v')
    return math.isclose(vds,.1,abs_tol=1e-12) and math.isclose(read,expected_read,abs_tol=1e-12)


def _pulse(prepared, result):
    c=result['settings']; states=[]
    for d, rows in prepared:
        if d.get('direction') not in ('ltp','ltd'):
            raise ValueError('Pulse direction must explicitly be ltp or ltd')
        if not _bias(d):
            raise ValueError('Pulse states require VDS=0.1 V and read VGS=0 V')
        read=[abs(r['vgs_v']) <= c['read_tolerance_v'] for r in rows]
        i=0; extraction_index=0
        while i < len(rows):
            if read[i]:
                i+=1; continue
            j=i
            while i < len(rows) and not read[i]: i+=1
            segment=rows[j:i]
            peaks=[r for r in segment if abs(r['vgs_v']) >= c['write_threshold_v']]
            if not peaks:
                _excluded(result,d,'write_threshold_not_reached',transition_row=rows[j]['source_row'])
                continue
            expected=-1 if d['direction']=='ltp' else 1
            if any(r['vgs_v']*expected <= 0 for r in segment):
                raise ValueError(f'Wrong pulse polarity in {d["file_id"]} at row {rows[j]["source_row"]}')
            if j < 2 or not all(read[j-2:j]) or rows[j]['source_row']-rows[j-2]['source_row'] != 2:
                _excluded(result,d,'insufficient_preceding_read_samples',transition_row=rows[j]['source_row']); continue
            r=rows[j-2]
            if r['time_s'] < c['start_time_s']:
                _excluded(result,d,'before_start_time',source_row=r['source_row'],transition_row=rows[j]['source_row']); continue
            extraction_index+=1
            identity=json.dumps([d['sha256'],d.get('sheet'),r['source_row'],PARSER_VERSION],separators=(',',':'),ensure_ascii=False)
            conductance=r['id_a']/c['vds_v']
            state=dict(state_id=hashlib.sha256(identity.encode()).hexdigest(),source_id=d['file_id'],source_row=r['source_row'],transition_row=rows[j]['source_row'],time_s=r['time_s'],direction=d['direction'],pulse_step=None,extraction_index=extraction_index,id_a=r['id_a'],vgs_v=r['vgs_v'],conductance_s=conductance,selected=conductance>0,exclusion_reason=None if conductance>0 else 'nonpositive_conductance')
            states.append(state)
            if conductance <= 0: _excluded(result,d,'nonpositive_conductance',source_row=r['source_row'])
        if read[-1]:
            start=len(rows)-1
            while start>0 and read[start-1]: start-=1
            _excluded(result,d,'no_next_transition',source_row=rows[start]['source_row'],end_row=rows[-1]['source_row'])
    if len({s['state_id'] for s in states}) != len(states):
        raise ValueError('The same physical source state was supplied more than once')
    result['states']=states; result['tables']['states']=deepcopy(states)
    result['summaries']={'candidate_count':len(states),'positive_candidate_count':sum(s['selected'] for s in states),'pools':make_pools(states)}


def make_pools(states):
    valid=[s for s in states if s['selected'] and s['conductance_s']>0]
    directions={d:[s for s in valid if s['direction']==d] for d in ('ltp','ltd')}
    lo=hi=None
    if all(directions.values()):
        lo=max(min(s['conductance_s'] for s in group) for group in directions.values())
        hi=min(max(s['conductance_s'] for s in group) for group in directions.values())
    common=[s for s in valid if lo is not None and lo<=s['conductance_s']<=hi] if lo is not None and lo<hi else []
    pools={}
    for name,group in {**directions,'combined':valid,'common':common}.items():
        values={s['conductance_s'] for s in group}; available=len(values)>=2
        pools[name]=dict(state_ids=[s['state_id'] for s in group],available=available,reason=None if available else ('no_common_range' if name=='common' and (lo is None or lo>=hi) else 'fewer_than_two_unique_states'),g_min_s=min(values) if values else None,g_max_s=max(values) if values else None,common_lo_s=lo if name=='common' else None,common_hi_s=hi if name=='common' else None)
    return pools


def _crossing(rows, xkey, ykey, target, selection):
    if selection:
        selected=selection.get('source_rows')
        if not selection.get('reason') or not isinstance(selected,list) or len(selected)!=2:
            raise ValueError('Crossing selection requires two adjacent source_rows and a reason')
        matches=[i for i in range(len(rows)-1) if [rows[i]['source_row'],rows[i+1]['source_row']]==selected]
        if not matches or selected[1]-selected[0]!=1:
            raise ValueError('Selected crossing rows must be adjacent original samples')
        rows=rows[matches[0]:matches[0]+2]
    values=[]; plateau=False
    for i,row in enumerate(rows):
        if row[xkey]==target: values.append((row[ykey],[row['source_row']]))
        if i+1 == len(rows): continue
        nxt=rows[i+1]
        if nxt['source_row']!=row['source_row']+1: continue
        x1,x2=row[xkey],nxt[xkey]
        if x1==x2==target: plateau=True
        elif min(x1,x2)<target<max(x1,x2):
            values.append((row[ykey]+(target-x1)*(nxt[ykey]-row[ykey])/(x2-x1),[row['source_row'],nxt['source_row']]))
    if plateau or len(values)>1:
        return dict(status='ambiguous_crossing',value=None,crossings=[dict(value=v,source_rows=r) for v,r in values],selection=selection)
    if not values:
        return dict(status='no_crossing',value=None,crossings=[],selection=selection)
    return dict(status='ok',value=values[0][0],crossings=[dict(value=values[0][0],source_rows=values[0][1])],selection=selection)


def _iv_meta(d, rows):
    if d.get('branch') not in ('program','erase'):
        raise ValueError('An explicit program/erase branch is required')
    amplitude=_finite(d.get('sweep_amplitude_v'),'sweep_amplitude_v')
    if not 1 <= amplitude <= 15:
        raise ValueError('Sweep amplitude must be between 1 and 15 V')
    if not math.isclose(_finite(d.get('vds_v'),'vds_v'),.1,abs_tol=1e-12):
        raise ValueError('IV analysis requires VDS=0.1 V')
    sign=1 if d['branch']=='erase' else -1
    if any((b['vgs_v']-a['vgs_v'])*sign < 0 for a,b in zip(rows,rows[1:])):
        raise ValueError('Branch direction conflicts with acquisition order; explicitly separate sweeps')
    return amplitude


def _stats(values):
    return (mean(values),stdev(values) if len(values)>1 else None) if values else (None,None)


def _iv(prepared,result):
    table=[]; pairs={}
    for d,rows in prepared:
        amplitude=_iv_meta(d,rows)
        key=(d['device_id'],amplitude,d['branch'])
        if key in pairs: raise ValueError('Duplicate device/amplitude/branch; choose one explicit sweep')
        crossing=_crossing(rows,'id_a','vgs_v',result['settings']['iref_a'],_crossing_selection(d,result))
        row=dict(file_id=d['file_id'],selection_key=crossing_selection_key(d),device_id=d['device_id'],sweep_amplitude_v=amplitude,branch=d['branch'],vth_v=crossing.pop('value'),max_id_a=max(r['id_a'] for r in rows),**crossing)
        table.append(row); pairs[key]=row
        if row['status']!='ok': _excluded(result,d,row['status'],branch=d['branch'],sweep_amplitude_v=amplitude)
    mw=[]
    for device,amplitude in sorted({(d,a) for d,a,_ in pairs}):
        p=pairs.get((device,amplitude,'program')); e=pairs.get((device,amplitude,'erase'))
        valid=p and e and p['status']==e['status']=='ok'
        mw.append(dict(device_id=device,sweep_amplitude_v=amplitude,mw_v=abs(p['vth_v']-e['vth_v']) if valid else None,status='ok' if valid else 'missing_valid_branch'))
    aggregates=[]
    for amplitude in sorted({row['sweep_amplitude_v'] for row in table}):
        entry=dict(sweep_amplitude_v=amplitude)
        for branch in ('program','erase'):
            values=[r['vth_v'] for r in table if r['sweep_amplitude_v']==amplitude and r['branch']==branch and r['status']=='ok']
            entry[branch+'_vth_mean_v'],entry[branch+'_vth_std_v']=_stats(values)
            entry[branch+'_n']=len(values)
            maxima=[r['max_id_a'] for r in table if r['sweep_amplitude_v']==amplitude and r['branch']==branch]
            entry[branch+'_max_id_mean_a']=mean(maxima) if maxima else None
        values=[r['mw_v'] for r in mw if r['sweep_amplitude_v']==amplitude and r['status']=='ok']
        entry['mw_mean_v'],entry['mw_std_v']=_stats(values); entry['mw_n']=len(values)
        aggregates.append(entry)
    result['tables'].update(vth=table,memory_window=mw,by_amplitude=aggregates)
    result['summaries']={'valid_vth_count':sum(r['status']=='ok' for r in table),'valid_mw_count':sum(r['status']=='ok' for r in mw)}


def _d2d(prepared,result):
    devices=sorted({d['device_id'] for d,_ in prepared})
    if len(devices)!=2: raise ValueError('D2D requires exactly two distinct physical devices')
    matched={}; table=[]
    for d,rows in prepared:
        # Bias mismatch is excluded explicitly; it is never silently pooled.
        if d.get('branch') not in ('program','erase'): raise ValueError('D2D requires an explicit branch')
        amp=_finite(d.get('sweep_amplitude_v'),'sweep_amplitude_v')
        if not 1<=amp<=15: raise ValueError('Sweep amplitude must be between 1 and 15 V')
        sign=1 if d['branch']=='erase' else -1
        if any((b['vgs_v']-a['vgs_v'])*sign<0 for a,b in zip(rows,rows[1:])): raise ValueError('Explicit branch conflicts with sweep direction')
        key=(amp,d['branch']); group=matched.setdefault(key,{})
        if d['device_id'] in group: raise ValueError('Duplicate physical device for a matched condition')
        crossing=_crossing(rows,'vgs_v','id_a',0.,_crossing_selection(d,result))
        current=crossing.pop('value'); g=current/.1 if current is not None else None
        reason=None if crossing['status']=='ok' else crossing['status']
        if not _bias(d): reason='read_condition_mismatch'
        elif g is not None and g<=0: reason='nonpositive_conductance'
        row=dict(file_id=d['file_id'],selection_key=crossing_selection_key(d),device_id=d['device_id'],sweep_amplitude_v=amp,branch=d['branch'],id_a=current,conductance_s=g,included=reason is None,reason=reason,**crossing)
        group[d['device_id']]=row
        if reason: _excluded(result,d,reason,sweep_amplitude_v=amp,branch=d['branch'])
    cvs=[]
    for (amp,branch),group in sorted(matched.items()):
        entries=[group.get(device) for device in devices]
        reason=None
        if any(e is None for e in entries): reason='missing_device_condition'
        elif any(not e['included'] for e in entries): reason='invalid_device_condition'
        key=f'{amp:g}:{branch}'
        if key in result['settings']['excluded_conditions']:
            note=result['settings']['excluded_conditions'][key]
            if not isinstance(note,str) or not note.strip(): raise ValueError('User exclusion requires a reason')
            reason='user_excluded: '+note
        gs=[e['conductance_s'] for e in entries if e and e['included']]
        avg=std=cv=None
        if reason is None:
            avg=mean(gs); std=stdev(gs); cv=std/avg; cvs.append(cv)
        else: result['exclusions'].append(dict(condition_key=key,reason=reason))
        table.append(dict(condition_key=key,sweep_amplitude_v=amp,branch=branch,device_values=entries,physical_device_count=2,mean_g_s=avg,std_g_s=std,cv=cv,included=reason is None,reason=reason))
    d2d=dict(status='available' if cvs else 'unavailable',cv=math.sqrt(mean([c*c for c in cvs])) if cvs else None,analysis_id=None,source_kind='iv_proxy',distribution='assumed_lognormal',physical_device_count=2,matched_conditions=len(cvs),assumption_ids=['d2d_lognormal','d2d_state_common'],reason=None if cvs else 'no_valid_matched_conditions')
    result['d2d']=d2d; result['tables']['d2d_conditions']=table; result['summaries']=deepcopy(d2d)
    result['warnings'].append('D2D lognormal distribution and state-common application are assumptions, not identified by two-device IV data')


def _retention(prepared,result):
    if len(prepared)!=1: raise ValueError('Select one explicit raw retention dataset per analysis')
    d,rows=prepared[0]
    if 'normal' in str(d.get('sheet','')).lower() or d.get('current_basis','raw') not in ('raw','absolute'):
        raise ValueError('Retention requires raw absolute current; normalized or offset data are invalid')
    expected=RETENTION_SOURCES.get(d['condition_id'])
    if expected and d.get('source_label') != expected[0]:
        raise ValueError(f'Retention source for {d["condition_id"]} must be {expected[0]}; R3(2) is excluded')
    if d.get('source_label')=='R3(2)': raise ValueError('R3(2) is explicitly excluded by the device team')
    expected_read=expected[1] if expected else _finite(d.get('read_vgs_v'),'read_vgs_v')
    if not _bias(d,expected_read): raise ValueError('Retention measurement bias does not match the approved source')
    kept=[]
    for r in rows:
        if r['time_s'] < result['settings']['retention_start_time_s']:
            _excluded(result,d,'before_fit_start',source_row=r['source_row'])
        else: kept.append(r)
    if len({r['time_s'] for r in kept})<3:
        raise ValueError('Retention OLS requires at least three distinct valid times >=10 s')
    x=[math.log10(r['time_s']) for r in kept]; xm=mean(x); xx=sum((v-xm)**2 for v in x)
    output=dict(status='available',analysis_id=None,source_label=d.get('source_label'),vds_v=.1,read_vgs_v=expected_read,current_basis='raw',reference_time_s=10.)
    fits=[]
    for direction in ('program','erase'):
        y=[r[direction+'_id_a'] for r in kept]; ym=mean(y)
        b=sum((xi-xm)*(yi-ym) for xi,yi in zip(x,y))/xx; a=ym-b*xm
        predicted=[a+b*v for v in x]; residual=[yi-pi for yi,pi in zip(y,predicted)]
        sse=sum(v*v for v in residual); sst=sum((yi-ym)**2 for yi in y)
        fit=dict(a=a,b=b,rmse=math.sqrt(sse/len(y)),r_squared=1-sse/sst if sst>0 else None,n=len(y),time_min_s=min(r['time_s'] for r in kept),time_max_s=max(r['time_s'] for r in kept),a_unit='A',b_unit='A/decade')
        output[direction+'_fit']=fit
        for r,pred,res in zip(kept,predicted,residual):
            fits.append(dict(file_id=d['file_id'],source_row=r['source_row'],direction=direction,time_s=r['time_s'],id_a=r[direction+'_id_a'],fit_id_a=pred,residual_a=res,instantaneous_slope_a_per_s=b/(r['time_s']*math.log(10))))
    output['program_reference_current_a']=output['program_fit']['a']+output['program_fit']['b']
    output['simulation_available']=output['program_reference_current_a']>0
    if not output['simulation_available']: result['warnings'].append('Program I_fit(10 s) is nonpositive; retention simulation is unavailable')
    if expected_read != 0: result['warnings'].append('Applying retention to pulse states at VGS=0 V assumes transfer across read bias')
    if d.get('header_read_vgs_v') is not None and d['header_read_vgs_v'] != expected_read:
        result['warnings'].append('Original header bias differs; the device-team Retention PPT bias takes precedence')
    result['retention']=output; result['tables']['retention_fit']=fits; result['summaries']=deepcopy(output)

