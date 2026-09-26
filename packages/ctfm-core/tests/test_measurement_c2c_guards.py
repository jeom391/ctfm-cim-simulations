"""08A guards: measurement-condition validation, real worksheet provenance, strict JSON, CLI behaviour."""
import hashlib
import importlib.util
import io
import json
import subprocess
import sys
from pathlib import Path

import numpy as np
import pytest
from openpyxl import Workbook

from ctfm.measurement import parse_table
from ctfm.measurement.c2c import C2CAnalysisError, analyze_c2c, analyze_c2c_file

CLI_PATH = Path(__file__).resolve().parents[3] / 'scripts/analyze_c2c.py'


def trend(n, coefficients, n_min, n_max):
    x = 2 * (n - n_min) / (n_max - n_min) - 1
    return sum(c * x ** k for k, c in enumerate(coefficients))


def workbook(cycles, program, erase, sheet='Origin_data'):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(['Cycle', 'Program_Id_uA', 'Erase_Id_uA', 'Erase_minus_Program_uA'])
    for c, p, e in zip(cycles, program, erase):
        ws.append([c, p, e, e - p])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def synthetic(n, seed=7, noise=0.01):
    cycles = np.arange(1, n + 1)
    rng = np.random.default_rng(seed)
    program = trend(cycles, [75, 3, -2, 1], 1, n) * (1 + noise * rng.standard_normal(n))
    erase = trend(cycles, [92, 2, -1.5, .5], 1, n) * (1 + noise * rng.standard_normal(n))
    return cycles.tolist(), program.tolist(), erase.tolist()


def book(n=40, **kw):
    return workbook(*synthetic(n), **kw)


def analyze(data, **kw):
    return analyze_c2c_file(data, 'synthetic.xlsx', **kw)


def codes(fn):
    with pytest.raises(C2CAnalysisError) as e:
        fn()
    return [i['code'] for i in e.value.issues]


def blocked_table():
    cycles = np.arange(1, 41)
    return parse_table(workbook(cycles.tolist(), np.linspace(1.0, -3.0, 40).tolist(),
                                trend(cycles, [90, 1, 0, 0], 1, 40).tolist()), 'x.xlsx')


def short_table():
    return parse_table(workbook(list(range(1, 6)), [70, 71, 70.5, 71.5, 72], [90, 91, 90.2, 91.4, 92]), 'x.xlsx')


@pytest.mark.parametrize('field,value,code', [
    ('vds_v', None, 'empty_condition_value'), ('program_voltage_v', None, 'empty_condition_value'),
    ('read_terminal_meaning', '', 'empty_condition_value'), ('read_extraction_point', '   ', 'empty_condition_value'),
    ('read_terminal_meaning', None, 'invalid_condition_type'), ('read_extraction_point', 5, 'invalid_condition_type'),
    ('vds_v', True, 'invalid_condition_type'), ('read_time_s', False, 'invalid_condition_type'),
    ('program_voltage_v', '10', 'invalid_condition_type'), ('erase_pulse_width_s', '0.01', 'invalid_condition_type'),
    ('vds_v', [0.1], 'invalid_condition_type'),
    ('program_pulse_width_s', float('nan'), 'non_finite_condition'), ('erase_voltage_v', float('inf'), 'non_finite_condition'),
    ('read_voltage_v', float('-inf'), 'non_finite_condition'), ('vds_v', float('nan'), 'non_finite_condition'),
    ('program_pulse_width_s', 0, 'non_positive_condition'), ('erase_pulse_width_s', -0.01, 'non_positive_condition'),
    ('read_time_s', 0.0, 'non_positive_condition'),
])
def test_invalid_measurement_conditions_are_rejected_never_confirmed(field, value, code):
    assert codes(lambda: analyze(book(), measurement_conditions={field: value})) == [code]


def test_valid_conditions_including_signs_zero_read_and_text_are_recorded_as_confirmed():
    r = analyze(book(), measurement_conditions=dict(
        program_voltage_v=10, program_pulse_width_s=0.01, erase_voltage_v=-10, erase_pulse_width_s=np.float64(0.01),
        read_voltage_v=0, read_terminal_meaning='  gate bias, drain floating ', vds_v=0.1, read_time_s=1e-3,
        read_extraction_point='end of read pulse'))
    c = r['measurement_conditions']
    assert c['program_voltage_v'] == dict(value=10.0, confirmed=True)
    assert c['erase_voltage_v'] == dict(value=-10.0, confirmed=True)
    assert c['read_voltage_v'] == dict(value=0.0, confirmed=True)
    assert c['read_terminal_meaning'] == dict(value='gate bias, drain floating', confirmed=True)
    assert all(v['confirmed'] for v in c.values())
    assert r['conductance_conversion'] == dict(available=True, reason=None)
    assert not any('not confirmed' in w for w in r['warnings'])
    json.dumps(r, allow_nan=False)


def test_vds_zero_is_recorded_but_conductance_conversion_is_declared_unavailable():
    r = analyze(book(), measurement_conditions=dict(vds_v=0))
    assert r['measurement_conditions']['vds_v'] == dict(value=0.0, confirmed=True)
    assert r['conductance_conversion']['available'] is False and 'undefined' in r['conductance_conversion']['reason']
    assert any('vds_v = 0' in w for w in r['warnings'])
    assert analyze(book())['conductance_conversion'] == dict(available=False, reason='vds_v is not confirmed')


def test_multiple_condition_problems_are_reported_together_and_non_object_is_rejected():
    got = codes(lambda: analyze(book(), measurement_conditions=dict(vds_v='x', read_time_s=-1, nope=1)))
    assert sorted(got) == ['invalid_condition_type', 'non_positive_condition', 'unknown_condition_field']
    assert codes(lambda: analyze(book(), measurement_conditions=[('vds_v', 0.1)])) == ['invalid_condition_type']


def two_sheet_book():
    wb = Workbook()
    wb.active.title = 'notes'
    wb.active.append(['just', 'notes'])
    data = wb.create_sheet('Origin_data')
    data.append(['Cycle', 'Program_Id_uA', 'Erase_Id_uA'])
    for n in range(1, 21):
        data.append([n, 70 + n / 10, 90 + n / 9])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_actual_worksheet_is_recorded_for_auto_explicit_multi_and_csv():
    auto = analyze(book())
    assert (auto['provenance']['sheet'], auto['provenance']['sheet_selection']) == ('Origin_data', 'auto_single_sheet')
    assert analyze(book(sheet='Data 2'))['provenance']['sheet'] == 'Data 2'
    explicit = analyze(book(), sheet='Origin_data')
    assert (explicit['provenance']['sheet'], explicit['provenance']['sheet_selection']) == ('Origin_data', 'explicit')
    assert codes(lambda: analyze(two_sheet_book())) == ['sheet_not_selected']
    picked = analyze(two_sheet_book(), sheet='Origin_data')
    assert (picked['provenance']['sheet'], picked['provenance']['sheet_selection']) == ('Origin_data', 'explicit')
    with pytest.raises(ValueError):
        analyze(book(), sheet='missing')
    csv = 'Cycle,Program_Id_uA,Erase_Id_uA\n' + ''.join(f'{n},{70 + n / 10},{90 + n / 9}\n' for n in range(1, 21))
    r = analyze_c2c_file(csv.encode(), 'x.csv')
    assert (r['provenance']['sheet'], r['provenance']['sheet_selection']) == (None, 'not_applicable_csv')
    assert codes(lambda: analyze_c2c_file(csv.encode(), 'x.csv', sheet='Origin_data')) == ['sheet_not_applicable']


def test_every_result_is_strict_json_including_blocked_short_and_null_diagnostics():
    results = [analyze(book()), analyze_c2c(blocked_table(), filename='x'), analyze_c2c(short_table(), filename='x')]
    for r in results:
        json.loads(json.dumps(r, allow_nan=False), parse_constant=lambda c: pytest.fail(f'non-standard JSON constant {c}'))
    assert results[1]['program']['status'] == 'blocked'


def test_overflowing_magnitudes_are_rejected_instead_of_emitting_infinity():
    cycles, program, erase = synthetic(30)
    huge = [v * 1e200 for v in program]
    assert codes(lambda: analyze(workbook(cycles, huge, erase))) == ['non_finite_result']


def run_cli(monkeypatch, *args):
    spec = importlib.util.spec_from_file_location('analyze_c2c_cli', CLI_PATH)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    monkeypatch.setattr(sys, 'argv', ['analyze_c2c.py', *args])
    return module.main()


def test_cli_writes_strict_json_plot_and_actual_sheet_without_touching_the_source(tmp_path, monkeypatch, capsys):
    src = tmp_path / 'c2c.xlsx'
    src.write_bytes(book(60))
    digest = hashlib.sha256(src.read_bytes()).hexdigest()
    out, png = tmp_path / 'r.json', tmp_path / 'r.png'
    code = run_cli(monkeypatch, '--file', str(src), '--out', str(out), '--plot', str(png), '--condition-id', 'A3',
                   '--condition', 'program_voltage_v=10', '--condition', 'erase_voltage_v=-10',
                   '--condition', 'read_voltage_v=0', '--condition', 'read_terminal_meaning=gate')
    assert code == 0 and png.stat().st_size > 1000
    r = json.loads(out.read_text(encoding='utf-8'), parse_constant=lambda c: pytest.fail(c))
    assert r['provenance']['sheet'] == 'Origin_data' and r['provenance']['condition_id'] == 'A3'
    assert r['measurement_conditions']['erase_voltage_v'] == dict(value=-10.0, confirmed=True)
    assert r['measurement_conditions']['vds_v']['confirmed'] is False
    assert 'sheet: Origin_data (auto_single_sheet)' in capsys.readouterr().out
    assert hashlib.sha256(src.read_bytes()).hexdigest() == digest


def test_cli_plots_blocked_and_short_inputs_without_crashing(tmp_path, monkeypatch):
    cases = {'blocked': workbook(list(range(1, 41)), np.linspace(1.0, -3.0, 40).tolist(),
                                 trend(np.arange(1, 41), [90, 1, 0, 0], 1, 40).tolist()),
             'short': workbook(list(range(1, 6)), [70, 71, 70.5, 71.5, 72], [90, 91, 90.2, 91.4, 92])}
    for name, data in cases.items():
        src = tmp_path / f'{name}.xlsx'
        src.write_bytes(data)
        png, out = tmp_path / f'{name}.png', tmp_path / f'{name}.json'
        assert run_cli(monkeypatch, '--file', str(src), '--plot', str(png), '--out', str(out), '--no-series') == 0
        assert png.stat().st_size > 1000 and 'series' not in json.loads(out.read_text(encoding='utf-8'))


def test_cli_rejects_bad_conditions_and_inputs_with_nonzero_exit(tmp_path, monkeypatch, capsys):
    src = tmp_path / 'c2c.xlsx'
    src.write_bytes(book(30))
    for bad in ('vds_v=nan', 'read_time_s=0', 'program_pulse_width_s=-1', 'erase_voltage_v=inf', 'read_extraction_point='):
        out = tmp_path / 'never.json'
        assert run_cli(monkeypatch, '--file', str(src), '--out', str(out), '--condition', bad) == 2, bad
        assert not out.exists()
        assert json.loads(capsys.readouterr().err)['error'] == 'invalid_c2c_input'
    for bad in ('vds_v=abc', 'bogus=1', 'vds_v'):
        with pytest.raises(SystemExit) as e:
            run_cli(monkeypatch, '--file', str(src), '--condition', bad)
        assert e.value.code not in (0, None)
    garbage = tmp_path / 'bad.xlsx'
    garbage.write_bytes(b'not a workbook')
    assert run_cli(monkeypatch, '--file', str(garbage)) == 2
    assert json.loads(capsys.readouterr().err)['error'] == 'unreadable_input'
    assert run_cli(monkeypatch, '--file', str(tmp_path / 'missing.xlsx')) == 2


def test_cli_process_exit_code_is_nonzero_for_invalid_condition(tmp_path):
    src = tmp_path / 'c2c.xlsx'
    src.write_bytes(book(30))
    bad = subprocess.run([sys.executable, str(CLI_PATH), '--file', str(src), '--condition', 'vds_v=nan'],
                         capture_output=True, text=True)
    assert bad.returncode == 2 and 'non_finite_condition' in bad.stderr
    ok = subprocess.run([sys.executable, str(CLI_PATH), '--file', str(src)], capture_output=True, text=True)
    assert ok.returncode == 0 and 'sheet: Origin_data' in ok.stdout
