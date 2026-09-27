"""Read every real IV Sweep / Retention workbook with the explicit-selection layout readers and report support.

    python scripts/verify_measured_layouts.py --root "<관련 자료>" --out report.json

The selections below are a *verification* choice (largest valid amplitude block, widest increasing segment = Erase,
widest decreasing segment = Program, the four Retention Raw Data columns in header order). They are recorded in the report and are not an
approval of any device or block for simulation. Source files are only read.
"""
import argparse
import hashlib
import json
import re
import sys
from pathlib import Path

from ctfm.measurement import RETENTION_SOURCES, analyze
from ctfm.measurement.layouts import iv_dataset, read_iv_blocks, read_retention_layout, retention_dataset

RETENTION_READ_VGS = {'A1': 0., 'A2': 0., 'A3': .5, 'A4': .1, 'A5': .5}
IV_UNITS = dict(vgs_v='V', id_a='A')


def sha(data):
    return hashlib.sha256(data).hexdigest()


def team_table(root):
    from openpyxl import load_workbook
    wb = load_workbook(root / 'Vth_CCM' / 'All Device Vth_CCM.xlsx', read_only=True, data_only=True)
    table = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=3, max_col=2, values_only=True):
            if row[0] is not None and isinstance(row[1], (int, float)):
                m = re.match(r'\d+', str(row[0]))
                if m:
                    table.setdefault(ws.title.strip(), {}).setdefault(m.group(), []).append(row[1])
    return table


def iv_report(root, team):
    out = []
    for path in sorted((root / 'IV Sweep').glob('A*/*.xlsx')):
        if path.name.startswith('~$'):
            continue
        cond = path.parent.name
        entry = dict(file=f'{cond}/{path.name}', condition=cond)
        data = path.read_bytes()
        entry['sha256'] = sha(data)
        try:
            layout = read_iv_blocks(data, path.name)
            entry.update(status='read', blocks=layout['block_count'], sheet=layout['sheet'],
                         points_per_block=sorted({b['points'] for b in layout['blocks'] if b['status'] == 'ok'}),
                         amplitudes=[b.get('proposed_amplitude_v') for b in layout['blocks']])
            entry['invalid_blocks'] = [dict(index=b['index'], error=b['error']) for b in layout['blocks'] if b['status'] != 'ok']
            good = [b for b in layout['blocks'] if b['status'] == 'ok']
            block = max(good, key=lambda b: b['proposed_amplitude_v'])  # largest amplitude present (15 V, or 14 V for A4 #238)
            entry['selected_amplitude_v'] = block['proposed_amplitude_v']
            segs = block['segments']
            entry['segments'] = [(s['direction'], s['vg_start'], s['vg_end']) for s in segs]
            span = lambda s: abs(s['vg_end'] - s['vg_start'])
            up_seg = max((s for s in segs if s['direction'] == 'increasing'), key=lambda s: (span(s), s['index']))
            down_seg = max((s for s in segs if s['direction'] == 'decreasing'), key=lambda s: (span(s), s['index']))
            entry['selected_segments'] = dict(erase=up_seg['index'], program=down_seg['index'])
            base = dict(sheet=layout['sheet'], file_id=path.name, sha256=entry['sha256'], device_id=path.stem,
                        condition_id=cond, units=IV_UNITS, block=block['index'], sweep_amplitude_v=block['proposed_amplitude_v'])
            up = iv_dataset(data, path.name, segment=up_seg['index'], branch='erase', **base)
            down = iv_dataset(data, path.name, segment=down_seg['index'], branch='program', **base)
            result = analyze('iv', [up, down], {})
            rows = {r['branch']: r for r in result['tables']['vth']}
            entry['erase_vth_v'] = rows['erase']['vth_v']; entry['erase_status'] = rows['erase']['status']
            entry['program_vth_v'] = rows['program']['vth_v']; entry['program_status'] = rows['program']['status']
            entry['memory_window_v'] = result['tables']['memory_window'][0]['mw_v']
            number = re.search(r'sweep_(\d+)', path.name).group(1)
            entry['team_vth_v'] = team.get(cond, {}).get(number)
            if entry['team_vth_v'] and entry['erase_vth_v'] is not None:
                entry['abs_diff_vs_team_v'] = [abs(entry['erase_vth_v'] - t) for t in entry['team_vth_v']]
        except Exception as exc:  # report, never hide
            entry.update(status='error', error=f'{type(exc).__name__}: {exc}')
        out.append(entry)
    return out


def retention_report(root):
    out = []
    for path in sorted((root / 'Retention').glob('_26CMFM_A*_Retention.xlsx')):
        cond = re.search(r'_A(\d)_', path.name).group(0).strip('_')
        entry = dict(file=path.name, condition=cond)
        data = path.read_bytes()
        entry['sha256'] = sha(data)
        try:
            layout = read_retention_layout(data, path.name)
            entry['columns'] = [c['header'] for c in layout['columns']]
            entry['embedded_source_headers'] = layout['embedded_source_headers']
            headers_bias = {tuple(e['read_bias_v']) for e in layout['embedded_source_headers']}
            header_bias = next(iter(headers_bias))[0] if len(headers_bias) == 1 and next(iter(headers_bias)) else None
            label = RETENTION_SOURCES[cond][0]
            dataset = retention_dataset(
                data, path.name, columns=dict(erase_time_s=0, erase_id_a=1, program_time_s=2, program_id_a=3),
                condition_id=cond, device_id=cond + '-retention', file_id=path.name, sha256=entry['sha256'], source_label=label,
                units=dict(erase_time_s='s', erase_id_a='A', program_time_s='s', program_id_a='A'),
                read_vgs_v=RETENTION_READ_VGS[cond], header_read_vgs_v=header_bias)
            result = analyze('retention', [dataset], {})
            entry.update(status='read', source_label=label, rows=len(dataset['rows']), skipped_blank_rows=dataset['skipped_blank_rows'],
                         excluded_before_10s=len(result['exclusions']),
                         program_fit=result['retention']['program_fit'], erase_fit=result['retention']['erase_fit'],
                         program_reference_current_a=result['retention']['program_reference_current_a'],
                         simulation_available=result['retention']['simulation_available'], warnings=result['warnings'])
        except Exception as exc:
            entry.update(status='error', error=f'{type(exc).__name__}: {exc}')
        out.append(entry)
    return out


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', required=True)
    parser.add_argument('--out', required=True)
    args = parser.parse_args()
    root = Path(args.root)
    report = dict(selection_note='verification selection only: largest-amplitude valid block; Erase = widest (last) increasing segment, Program = widest (last) decreasing segment; not an approval of any device/block for simulation',
                  iv=iv_report(root, team_table(root)), retention=retention_report(root))
    Path(args.out).write_text(json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')
    bad = [e['file'] for e in report['iv'] + report['retention'] if e['status'] != 'read']
    print(f"iv files: {len(report['iv'])}, retention files: {len(report['retention'])}, failed: {bad}")
    return 1 if bad else 0


if __name__ == '__main__':
    sys.exit(main())
